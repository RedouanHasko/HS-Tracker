from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


BOOK = Path(r"D:\Download\Location dates.xlsx")


def style_signature(cell):
    fill = cell.fill.fgColor.rgb or cell.fill.fgColor.indexed or ""
    border = (
        cell.border.left.style,
        cell.border.right.style,
        cell.border.top.style,
        cell.border.bottom.style,
    )
    return cell.style_id, str(fill), border


workbook = load_workbook(BOOK, data_only=False)
for sheet_name in ("Feuil1", "2025", "2026"):
    sheet = workbook[sheet_name]
    print(f"\nSHEET {sheet_name}")
    for row in range(4, sheet.max_row + 1):
        property_code = sheet.cell(row, 2).value
        if not property_code:
            continue
        print(f"ROW {row} PROPERTY {property_code}")
        start = 3
        while start <= 33:
            signature = style_signature(sheet.cell(row, start))
            end = start
            while end + 1 <= 33 and style_signature(sheet.cell(row, end + 1)) == signature:
                end += 1
            values = [
                f"{get_column_letter(column)}={sheet.cell(row, column).value!r}"
                for column in range(start, end + 1)
                if sheet.cell(row, column).value not in (None, "")
            ]
            if signature[0] not in (0, 5, 57) or values:
                print(
                    f"  days {start - 2:02d}-{end - 2:02d} "
                    f"cols {get_column_letter(start)}:{get_column_letter(end)} "
                    f"style={signature[0]} fill={signature[1]} border={signature[2]} "
                    f"values={', '.join(values) or '-'}"
                )
            start = end + 1
